#!/usr/bin/env python3
"""The voice-line spreadsheet for recording: both maps' full transcripts, and one
tab per character with each of their lines once, in reading order.
Usage: node tools/voice/transcript.mjs > lines.json
       python3 tools/voice/sheet.py . lines.json "voice lines.xlsx"
"""
import csv, json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPO, LINES, OUT = sys.argv[1:4]
rows = json.load(open(LINES))
script = {(r['filename'].split('_')[0], r['text']): r['filename']
          for r in csv.DictReader(open(f'{REPO}/tools/voice/script.csv', encoding='utf-8'))}
# lines that already have a recording in the game
vl = open(f'{REPO}/src/lore/voiceLines.js', encoding='utf-8').read()
recorded, who = set(), None
for line in vl.splitlines():
  m = re.match(r'  (\w+): \{', line)
  if m: who = m.group(1); continue
  m = re.match(r'    (".*"): \[', line)
  if m and who: recorded.add((who, json.loads(m.group(1))))

import os
if os.path.exists(f'{REPO}/assets/voice/cheddar/wanna-play-2k.mp3'): recorded.add(('cheddar', 'Wanna play 2K?'))
NAME = {'erik': 'Erik', 'kearns': 'Kearns', 'ryan': 'Ryan', 'pit': 'Pit', 'rocco': 'Rocco', 'brian': 'Brian',
        'schnitz': 'The Schnitz', 'zach': 'Zach', 'chops': 'Chops', 'regs': 'Regs', 'p': 'P', 'cheddar': 'The Cheddars'}
ORDER = list(NAME)
for r in rows:
  r['clip'] = script.get((r['who'], r['text']))
  assert r['clip'], r

F = 'Arial'
head_font = Font(name=F, bold=True, color='FFFFFF')
head_fill = PatternFill('solid', start_color='5A2E1A')
part_fill = PatternFill('solid', start_color='F3E3D3')
body = Font(name=F)
wrap = Alignment(wrap_text=True, vertical='top')
top = Alignment(vertical='top')
thin = Border(bottom=Side(style='thin', color='DDDDDD'))

def sheet(ws, headers, widths, data, part_col=None):
  ws.append(headers)
  for c in ws[1]: c.font, c.fill, c.alignment = head_font, head_fill, Alignment(vertical='center', wrap_text=True)
  ws.freeze_panes = 'A2'
  for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
  for d in data:
    ws.append(d)
    for c in ws[ws.max_row]:
      c.font, c.border = body, thin
      c.alignment = wrap if ws.column_dimensions[c.column_letter].width >= 30 else top
  ws.auto_filter.ref = ws.dimensions

wb = Workbook()
guide = wb.active
guide.title = 'How to use'

# --- full transcripts, one tab per map
first_at = {}
for mp in ['Stew Leonard High', 'Call of the Crust']:
  ws = wb.create_sheet(mp)
  data, n = [], 0
  for r in rows:
    if r['map'] != mp: continue
    n += 1
    k = (r['who'], r['text'])
    note = r['note']
    if k in first_at:
      note = '; '.join(x for x in [note, f"same line as {first_at[k]}, record once"] if x)
    else:
      first_at[k] = f"{'School' if mp.startswith('Stew') else 'Crust'} #{n}"
    data.append([n, r['part'], r['when'], NAME[r['who']], r['text'], note, r['clip']])
  sheet(ws, ['#', 'Part of the game', 'When it plays', 'Character', 'Line', 'Notes', 'Clip ID'],
        [6, 30, 34, 13, 70, 34, 20], data)

# --- one tab per character: each line once, in the order to read them
counts = []
for who in ORDER:
  seen, data = set(), []
  for r in rows:
    if r['who'] != who or r['text'] in seen: continue
    seen.add(r['text'])
    data.append([len(data) + 1, r['text'], 'School' if r['map'].startswith('Stew') else 'Crust', r['when'], r['note'],
                 'Yes' if (who, r['text']) in recorded else 'No', r['clip']])
  if not data: continue
  ws = wb.create_sheet(NAME[who].replace('The ', ''))
  sheet(ws, ['#', 'Line', 'Map', 'When it plays', 'Notes', 'Already in game?', 'Clip ID'],
        [6, 70, 9, 34, 30, 11, 20], data)
  counts.append((NAME[who], ws.title))

# --- the guide
g = guide
g.column_dimensions['A'].width = 24
g.column_dimensions['B'].width = 16
g.column_dimensions['C'].width = 16
g.column_dimensions['D'].width = 16
g['A1'] = 'Stew Zombies 2.0: voice line scripts'
g['A1'].font = Font(name=F, bold=True, size=14)
notes = [
  'The "Stew Leonard High" and "Call of the Crust" tabs are the full transcripts: every line, in roughly the order a game plays them.',
  'Each character has their own tab with every line they say, listed once, across both maps. That is the script to record from.',
  'Record one audio file per character. Read the lines top to bottom in the tab\'s order and leave about 1 second of silence between lines, so the file can be cut back into lines.',
  'Lines that come up in more than one place (for example the intercom lines every character says) are only on each character\'s tab once.',
  '"Already in game?" = Yes means the game already has a recording of that line (from the earlier batch of recordings). Record it again if you want your own take on it.',
  'Clip ID is the name the game knows the line by. You do not need it; it is how the lines get wired in.',
  'Not included: the Stew song (it is a music track) and on-screen text like notes and prompts.',
]
for i, t in enumerate(notes):
  c = g.cell(row=3 + i, column=1, value=t)
  c.font, c.alignment = body, Alignment(wrap_text=True, vertical='top')
  g.merge_cells(start_row=3 + i, start_column=1, end_row=3 + i, end_column=6)
  g.row_dimensions[3 + i].height = 30
r0 = 4 + len(notes)
for j, h in enumerate(['Character', 'Lines to record', 'Already in game', 'Still needed']):
  c = g.cell(row=r0, column=1 + j, value=h)
  c.font, c.fill = head_font, head_fill
for i, (name, tab) in enumerate(counts):
  r = r0 + 1 + i
  g.cell(row=r, column=1, value=name)
  g.cell(row=r, column=2, value=f"=COUNTA('{tab}'!B:B)-1")
  g.cell(row=r, column=3, value=f"=COUNTIF('{tab}'!F:F,\"Yes\")")
  g.cell(row=r, column=4, value=f'=B{r}-C{r}')
  for c in g[r][:4]: c.font = body
rt = r0 + 1 + len(counts)
g.cell(row=rt, column=1, value='Total')
for col in 'BCD':
  g[f'{col}{rt}'] = f'=SUM({col}{r0 + 1}:{col}{rt - 1})'
for c in g[rt][:4]: c.font = Font(name=F, bold=True)
rs = rt + 2
g.cell(row=rs, column=1, value='Lines in the school transcript').font = body
g.cell(row=rs, column=2, value="=COUNTA('Stew Leonard High'!A:A)-1").font = body
g.cell(row=rs + 1, column=1, value='Lines in the Crust transcript').font = body
g.cell(row=rs + 1, column=2, value="=COUNTA('Call of the Crust'!A:A)-1").font = body
wb.save(OUT)
print('saved', OUT, counts)
